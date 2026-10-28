# ChromaDB needs a newer SQLite than some Linux hosts (like Streamlit Cloud)
# ship with. If pysqlite3-binary is installed, use it in place of the built-in
# sqlite3. On Windows/local runs it's simply not installed and this does nothing.
try:
    __import__("pysqlite3")
    import sys
    sys.modules["sqlite3"] = sys.modules.pop("pysqlite3")
except ImportError:
    pass

import os
import time
import hashlib
import hmac
import html as html_lib
import inspect
import io
import json
import logging
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
import random
import re
import secrets
from datetime import datetime
import importlib.util
import streamlit as st
from dotenv import load_dotenv
from pydantic import BaseModel, Field
from typing import List
from langchain_google_genai import GoogleGenerativeAIEmbeddings, ChatGoogleGenerativeAI
from langchain_chroma import Chroma
from langchain_core.messages import HumanMessage, SystemMessage

load_dotenv()

# On Streamlit Community Cloud, API keys come from the app's Secrets manager
# (st.secrets) rather than a local .env file. This bridges that into the
# environment so the rest of the code works unchanged whether running
# locally (.env) or deployed (Streamlit Cloud secrets).
try:
    if "GOOGLE_API_KEY" in st.secrets:
        os.environ["GOOGLE_API_KEY"] = st.secrets["GOOGLE_API_KEY"]
except Exception:
    pass  # No secrets.toml locally — that's fine, .env already covers local use

DOCS_PATH = "docs"
PERSIST_DIRECTORY = "db/chroma_db"
EMBEDDING_MODEL_NAME = "models/gemini-embedding-001"
CHAT_MODEL_NAME = "gemini-3.1-flash-lite"


def load_module_from_file(module_name, file_path):
    """Filenames starting with a digit (e.g. 1_ingestion_pipeline.py) can't use a
    normal `import` statement, so we load them dynamically instead. This keeps
    1_ingestion_pipeline.py and 3_answer_generation.py completely untouched."""
    spec = importlib.util.spec_from_file_location(module_name, file_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# Load the existing ingestion file as a reusable module (its functions only run when called)
ingestion = load_module_from_file("ingestion_pipeline", "1_ingestion_pipeline.py")


_CHAT_MODEL = None


def get_chat_model():
    """One shared Gemini chat model. timeout/max_retries are set low on purpose:
    the library default retries rate-limit errors silently up to 6 times with
    growing waits, which looks like the app is frozen."""
    global _CHAT_MODEL
    if _CHAT_MODEL is None:
        _CHAT_MODEL = ChatGoogleGenerativeAI(model=CHAT_MODEL_NAME, timeout=60, max_retries=2)
    return _CHAT_MODEL


@st.cache_resource
def get_embedding_model():
    return GoogleGenerativeAIEmbeddings(model=EMBEDDING_MODEL_NAME)


def load_existing_vectorstore():
    """Load the vector store from disk if it already exists and has data."""
    if not os.path.exists(PERSIST_DIRECTORY):
        return None

    vectorstore = Chroma(
        persist_directory=PERSIST_DIRECTORY,
        embedding_function=get_embedding_model(),
        collection_metadata={"hnsw:space": "cosine"}
    )
    if vectorstore._collection.count() > 0:
        return vectorstore
    return None


def describe_page_with_gemini(image_bytes):
    """Send a rendered PDF page image to Gemini's vision model and get back
    a full transcription: regular text, tables reconstructed as markdown,
    and descriptions of any figures/charts/diagrams."""
    import base64
    b64_image = base64.b64encode(image_bytes).decode("utf-8")

    model = get_chat_model()
    message = HumanMessage(content=[
        {
            "type": "text",
            "text": (
                "Transcribe everything on this page factually and completely:\n"
                "- All regular text, exactly as written\n"
                "- Any tables: reconstruct them as a markdown table with correct rows and columns\n"
                "- Any figures, charts, or diagrams: describe what they show in plain text\n"
                "No commentary, no summarizing — just the full factual content."
            )
        },
        {
            "type": "image_url",
            "image_url": {"url": f"data:image/png;base64,{b64_image}"}
        }
    ])
    result = model.invoke([message])

    return extract_text_from_gemini_response(result)


def extract_text_from_gemini_response(result):
    """Safely pull text out of a Gemini response, regardless of whether it
    comes back as a plain string, a list of content blocks, or an empty
    response (which can happen if a safety filter silently blocks content)."""
    content = result.content

    if isinstance(content, str):
        return content

    if isinstance(content, list):
        text_parts = []
        for block in content:
            if isinstance(block, dict) and block.get("type") == "text":
                text_parts.append(block.get("text", ""))
            elif isinstance(block, str):
                text_parts.append(block)
        if text_parts:
            return "\n".join(text_parts)
        return "[No text returned for this page — it may have been blocked by a safety filter or contains no readable content.]"

    return "[Unexpected response format from the model for this page.]"

    if isinstance(result.content, list):
        return result.content[0]['text']
    return result.content


def describe_page_with_gemini_with_retry(image_bytes, max_retries=4):
    """Same as describe_page_with_gemini, but retries with backoff on rate
    limits — important on the free tier, where a page-heavy PDF can easily
    burst past the per-minute limit for the chat/vision model."""
    for attempt in range(max_retries):
        try:
            return describe_page_with_gemini(image_bytes)
        except Exception as e:
            error_text = str(e)
            if ("429" in error_text or "RESOURCE_EXHAUSTED" in error_text) and attempt < max_retries - 1:
                wait_time = 5 * (attempt + 1)
                time.sleep(wait_time)
            else:
                raise


def load_pdf_with_vision(file_path, min_text_length=50, page_callback=None, max_workers=4):
    """Extract text from each PDF page normally. Pages with little text, or
    that contain images (often tables/figures), are rendered and sent to
    Gemini vision. Those vision calls now run up to `max_workers` at a time
    instead of one by one, and there is no fixed pause between pages
    (rate-limit errors are still handled by retry-with-backoff)."""
    import fitz  # PyMuPDF
    from langchain_core.documents import Document

    pdf = fitz.open(file_path)
    total = len(pdf)
    contents = {}
    vision_jobs = {}

    for page_num, page in enumerate(pdf, start=1):
        native_text = page.get_text().strip()
        has_images = len(page.get_images()) > 0
        if len(native_text) >= min_text_length and not has_images:
            contents[page_num] = native_text
        else:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
            vision_jobs[page_num] = pix.tobytes("png")
    pdf.close()

    done = len(contents)
    if page_callback:
        page_callback(done, total)

    if vision_jobs:
        get_chat_model()  # create the shared model once, in the main thread
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {
                executor.submit(describe_page_with_gemini_with_retry, image_bytes): num
                for num, image_bytes in vision_jobs.items()
            }
            for future in as_completed(futures):
                contents[futures[future]] = future.result()
                done += 1
                if page_callback:
                    page_callback(done, total)

    return [
        Document(page_content=contents[n], metadata={"source": file_path, "page": n})
        for n in sorted(contents)
    ]


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}


def load_image_as_document(file_path):
    """Load a standalone image file (not embedded in a PDF) by sending it
    straight to Gemini's vision model — the same transcription/description
    pipeline already used for image-heavy PDF pages, so a photo of a
    document, a screenshot, a chart, or a diagram all get read the same way
    and become fully searchable text."""
    from langchain_core.documents import Document

    with open(file_path, "rb") as f:
        image_bytes = f.read()

    content = describe_page_with_gemini_with_retry(image_bytes)

    return [Document(
        page_content=content,
        metadata={"source": file_path}
    )]


def load_document_by_extension(file_path, page_callback=None):
    """Load a document. Supports .txt, .docx, .pdf (with vision fallback for
    tables/images/scans), and standalone image files (.jpg/.png/etc). Tries
    several encodings for .txt since Windows-generated text files aren't
    always plain UTF-8."""
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".txt":
        from langchain_community.document_loaders import TextLoader

        encodings_to_try = ["utf-8", "utf-8-sig", "utf-16", "cp1252", "latin-1"]
        last_error = None

        for encoding in encodings_to_try:
            try:
                loader = TextLoader(file_path, encoding=encoding)
                return loader.load()
            except Exception as e:
                last_error = e
                continue

        raise ValueError(f"Could not read file with any known encoding. Last error: {last_error}")

    elif extension == ".docx":
        from langchain_community.document_loaders import Docx2txtLoader
        loader = Docx2txtLoader(file_path)
        return loader.load()

    elif extension == ".pdf":
        return load_pdf_with_vision(file_path, page_callback=page_callback)

    elif extension in IMAGE_EXTENSIONS:
        return load_image_as_document(file_path)

    else:
        raise ValueError(f"Unsupported file type: {extension}. Supported: .txt, .docx, .pdf, .jpg, .jpeg, .png, .webp, .bmp")


def stamp_user_id(chunks, user_id):
    """Tag every chunk's metadata with the uploading user's ID.

    This is the entire basis of multi-user isolation in this app: there is
    only ONE shared ChromaDB collection on disk (one `db/chroma_db` folder),
    used by every visitor. Without this tag — and without filtering every
    search/count/delete by it afterward — any user could retrieve, count,
    or even wipe any other user's documents, since Chroma has no built-in
    concept of "owner" on its own."""
    for chunk in chunks:
        chunk.metadata["user_id"] = user_id
    return chunks


def count_user_chunks(vectorstore, user_id):
    """Count only the chunks belonging to this user, not the whole shared
    collection. Chroma's `.count()` has no filter, so we fetch the matching
    ids (metadata only, no embeddings/text) and count those instead."""
    if vectorstore is None:
        return 0
    result = vectorstore._collection.get(where={"user_id": user_id}, include=[])
    return len(result["ids"])


def list_user_documents(vectorstore, user_id):
    """Return {filename: chunk_count} for every file this user has ingested,
    read from the database itself, so it survives page refreshes."""
    if vectorstore is None:
        return {}
    result = vectorstore._collection.get(where={"user_id": user_id}, include=["metadatas"])
    files = {}
    for metadata in result["metadatas"]:
        name = os.path.basename(str((metadata or {}).get("source", "unknown")).replace("\\", "/"))
        files[name] = files.get(name, 0) + 1
    return dict(sorted(files.items()))


def delete_user_document(vectorstore, user_id, filename):
    """Delete ONE file's chunks (only this user's) from the database, and the
    saved copy from disk. Matches on the file name, so it works regardless of
    whether the stored path used / or \\ separators."""
    result = vectorstore._collection.get(where={"user_id": user_id}, include=["metadatas"])
    ids_to_delete = [
        doc_id for doc_id, metadata in zip(result["ids"], result["metadatas"])
        if os.path.basename(str((metadata or {}).get("source", "")).replace("\\", "/")) == filename
    ]
    if ids_to_delete:
        vectorstore._collection.delete(ids=ids_to_delete)

    saved_copy = os.path.join(DOCS_PATH, user_id, filename)
    if os.path.exists(saved_copy):
        try:
            os.remove(saved_copy)
        except OSError:
            pass
    return len(ids_to_delete)


def ingest_uploaded_files(uploaded_files, user_id, progress_callback=None):
    """Save uploaded files into a per-user subfolder of docs/, then load and
    embed ONLY those new files. Every resulting chunk is tagged with user_id
    before embedding.

    Files are saved under docs/<user_id>/ rather than directly in docs/ —
    without this, two different users uploading a same-named file (e.g. both
    named "resume.pdf") would silently overwrite each other on disk before
    either one gets embedded."""
    user_docs_path = os.path.join(DOCS_PATH, user_id)
    os.makedirs(user_docs_path, exist_ok=True)

    if progress_callback:
        progress_callback(0, 0, status="Saving uploaded files...")

    new_file_paths = []
    for uploaded_file in uploaded_files:
        file_path = os.path.join(user_docs_path, uploaded_file.name)
        with open(file_path, "wb") as f:
            f.write(uploaded_file.getbuffer())
        new_file_paths.append(file_path)

    documents = []
    failed_files = []
    for path in new_file_paths:
        try:
            if progress_callback:
                progress_callback(0, 0, status=f"Reading {os.path.basename(path)}...")

                def page_callback(page_num, total_pages, _path=path):
                    progress_callback(page_num, total_pages, status=f"Reading {os.path.basename(_path)}: page {page_num}/{total_pages}...")
                documents.extend(load_document_by_extension(path, page_callback=page_callback))
            else:
                documents.extend(load_document_by_extension(path))
        except Exception as e:
            failed_files.append((os.path.basename(path), str(e)))

    # Failures are recorded for the UI (file cards show them) instead of
    # being printed as a one-off warning that disappears on the next rerun.
    if "file_errors" not in st.session_state:
        st.session_state.file_errors = {}
    for name, error in failed_files:
        st.session_state.file_errors[name] = error

    if not documents:
        return st.session_state.vectorstore

    if progress_callback:
        progress_callback(0, 0, status="Splitting text into chunks...")

    chunks = ingestion.split_documents(documents)

    if not chunks:
        failed_names = {name for name, _ in failed_files}
        for path in new_file_paths:
            name = os.path.basename(path)
            if name not in failed_names:
                st.session_state.file_errors[name] = (
                    "No extractable text found - it may be a scanned/image-only file or an empty document."
                )
        return st.session_state.vectorstore

    stamp_user_id(chunks, user_id)

    embedding_model = get_embedding_model()

    if progress_callback:
        progress_callback(0, len(chunks))

    vectorstore = embed_chunks_with_retry(
        chunks, embedding_model, st.session_state.vectorstore,
        progress_callback=progress_callback
    )

    if progress_callback:
        progress_callback(len(chunks), len(chunks))

    return vectorstore


def embed_chunks_with_retry(chunks, embedding_model, existing_vectorstore, batch_size=50, max_retries=5, progress_callback=None):
    """Embed chunks in small batches with automatic retry-with-backoff on rate
    limits. Even on a paid tier, a burst of many chunks at once (like a large
    image-heavy PDF that needed vision on every page) can hit a temporary
    rate limit — this recovers instead of crashing the whole upload."""
    vectorstore = existing_vectorstore
    total = len(chunks)

    for i in range(0, total, batch_size):
        batch = chunks[i:i + batch_size]

        if progress_callback:
            progress_callback(i, total, status=f"Embedding chunk {i + 1}-{min(i + batch_size, total)}/{total}...")

        for attempt in range(max_retries):
            try:
                if vectorstore is None:
                    vectorstore = Chroma.from_documents(
                        documents=batch,
                        embedding=embedding_model,
                        persist_directory=PERSIST_DIRECTORY,
                        collection_metadata={"hnsw:space": "cosine"}
                    )
                else:
                    vectorstore.add_documents(batch)
                break
            except Exception as e:
                error_text = str(e)
                if ("429" in error_text or "RESOURCE_EXHAUSTED" in error_text) and attempt < max_retries - 1:
                    wait_time = 5 * (attempt + 1)
                    if progress_callback:
                        progress_callback(i, total, status=f"Rate limited, waiting {wait_time}s before retrying...")
                    time.sleep(wait_time)
                else:
                    raise


    return vectorstore


class QueryVariations(BaseModel):
    """Alternative phrasings of the user's question, preserving its exact
    original intent — used to improve retrieval recall."""
    variations: List[str] = Field(
        description=(
            "3 to 4 alternative ways to ask the exact same question. "
            "Preserve the original meaning and intent precisely — do not "
            "introduce new topics, assumptions, or change what is being asked."
        )
    )


def generate_query_variations(query, max_retries=2, model=None):
    """Ask Gemini for a few alternative phrasings of the question. Different
    wording can match different chunks in the vector store that the user's
    exact phrasing alone might miss (e.g. "wage employment" vs "salaried job").

    This costs one extra Gemini call per question. If it fails for any reason
    (including a rate limit), we fall back to just the original query rather
    than blocking the whole answer on it."""
    model = model or get_chat_model()
    structured_model = model.with_structured_output(QueryVariations)

    prompt = (
        "Generate 3 to 4 alternative phrasings of this question. Preserve its "
        "exact original meaning and intent — do not change what is being "
        "asked or introduce new topics. If the question contains a specific "
        "name, person, file name, or proper noun, keep it EXACTLY as written "
        "in every variation — never replace it with a pronoun, a generic "
        "term, or drop it.\n\n"
        f"Question: {query}"
    )

    for attempt in range(max_retries):
        try:
            result = structured_model.invoke(prompt)
            return [v.strip() for v in result.variations if v.strip()][:4]
        except Exception as e:
            error_text = str(e)
            if ("429" in error_text or "RESOURCE_EXHAUSTED" in error_text) and attempt < max_retries - 1:
                time.sleep(3 * (attempt + 1))
            else:
                return []  # Fall back to just the original query
    return []


def _search_one(vectorstore, q, k, user_filter):
    try:
        return vectorstore.similarity_search_with_relevance_scores(q, k=k, filter=user_filter)
    except Exception:
        return [(doc, None) for doc in vectorstore.similarity_search(q, k=k, filter=user_filter)]


def retrieve_with_multi_query(vectorstore, query, user_id, k_per_query=6):
    """Search with the original question plus a few Gemini-generated
    rephrasings, then combine and deduplicate the results.

    Speed: the original question is searched WHILE the rephrasings are being
    generated, and all the rephrasing searches then run in parallel, instead
    of everything happening one step after another.

    Every search is filtered to this user's own chunks (filter={"user_id":
    user_id}), so one user's question can never retrieve another user's
    documents. Dedup key: chunk_id if present, else a content hash. When a
    chunk is returned by several queries, its best score is kept."""
    user_filter = {"user_id": user_id}
    model = get_chat_model()  # created in the main thread, reused by workers

    with ThreadPoolExecutor(max_workers=6) as executor:
        variations_future = executor.submit(generate_query_variations, query, 2, model)
        original_future = executor.submit(_search_one, vectorstore, query, k_per_query, user_filter)

        variations = variations_future.result()

        unique_queries = [query]
        seen = {query.strip().lower()}
        for v in variations:
            key = v.strip().lower()
            if key and key not in seen:
                seen.add(key)
                unique_queries.append(v)

        extra_futures = [
            executor.submit(_search_one, vectorstore, q, k_per_query, user_filter)
            for q in unique_queries[1:]
        ]
        all_results = [original_future.result()] + [f.result() for f in extra_futures]

    candidates = {}  # dedup_key -> (doc, best_score)
    for results in all_results:
        for doc, score in results:
            dedup_key = doc.metadata.get("chunk_id")
            if not dedup_key:
                dedup_key = hashlib.md5(doc.page_content.encode("utf-8")).hexdigest()

            if dedup_key not in candidates:
                candidates[dedup_key] = (doc, score)
            else:
                existing_doc, existing_score = candidates[dedup_key]
                if score is not None and (existing_score is None or score > existing_score):
                    candidates[dedup_key] = (doc, score)

    return list(candidates.values()), unique_queries


def rerank_with_bm25(query, candidates, top_n=8):
    """Rerank the candidate pool against the ORIGINAL user query using BM25
    keyword relevance (via rank_bm25 — a small, pure-Python library, no
    model download and no extra API calls).

    Why BM25 instead of a neural cross-encoder or an external rerank API:
    - It's genuinely lightweight: no torch/transformers install, no model
      download, installs in seconds.
    - No extra network calls, so it can't fail from a rate limit or add
      latency from an external API — important given how much this project
      has already had to work around free-tier limits.
    - It scores directly against the user's actual question (not the
      generated variations), which keeps the final answer grounded in what
      was really asked rather than diluted across multiple phrasings.
    A heavier cross-encoder model (e.g. sentence-transformers) or a hosted
    reranker (e.g. Cohere Rerank, which has a free trial tier with its own
    rate limits) would likely rerank slightly more accurately, but at real
    cost in install size, latency, or external dependency — BM25 is the
    practical choice for this app's scale."""
    from rank_bm25 import BM25Okapi

    if not candidates:
        return []

    docs = [doc for doc, _ in candidates]
    tokenized_corpus = [doc.page_content.lower().split() for doc in docs]
    bm25 = BM25Okapi(tokenized_corpus)

    tokenized_query = query.lower().split()
    scores = bm25.get_scores(tokenized_query)

    scored = list(zip(docs, scores))
    scored.sort(key=lambda pair: pair[1], reverse=True)

    return [doc for doc, _ in scored[:top_n]]


def condense_question_with_history(query, chat_history, max_turns=3, max_retries=3):
    """Rewrite a follow-up question into a standalone one, using recent chat
    history for context — e.g. "what about the second one?" becomes "what
    is the second essential quality of a successful entrepreneur listed by
    Dr. Kiran Mazumdar-Shaw?" based on the previous exchange.

    Only the last few turns are used (max_turns) to keep the prompt small —
    older context matters less for resolving "it"/"that"/"the second one"
    style references, and keeping history short avoids needlessly growing
    the token count (and therefore cost/latency) with every message.

    If there's no history yet, or the rewrite fails for any reason
    (including a rate limit), we fall back to the original question
    unchanged — this should never be the reason a question fails."""
    if not chat_history:
        return query

    # Skip the extra Gemini call when the question clearly stands on its own:
    # long, and with no "it/that/the second one"-style reference.
    vague_words = {"it", "its", "that", "this", "these", "those", "they", "them", "he", "she",
                   "him", "her", "his", "their", "above", "previous", "earlier", "former",
                   "latter", "first", "second", "third", "last", "same"}
    words = re.findall(r"[a-z']+", query.lower())
    if len(words) > 6 and not (vague_words & set(words)):
        return query

    recent = chat_history[-(max_turns * 2):]
    history_text = "\n".join(f"{m['role'].capitalize()}: {m['content']}" for m in recent)

    model = get_chat_model()

    prompt = f"""Given this recent conversation and a follow-up question, decide whether the follow-up genuinely depends on the conversation to be understood, or whether it introduces a new, different subject.

Rewrite ONLY if the follow-up question contains a vague reference that depends on the immediately preceding exchange to resolve — things like "it", "that", "this", "the second one", "him", "her", "the document" with no name attached. In that case, replace the vague reference with what it refers to.

Do NOT rewrite, and return the follow-up completely UNCHANGED, if:
- It mentions a specific name, document, file, or topic not already the exact subject of the last exchange (a topic change)
- It's already clear and understandable on its own
- You are not confident what the vague reference points to

When in doubt, prefer returning the question UNCHANGED — incorrectly carrying over old context to a new topic is worse than occasionally missing a legitimate follow-up.

Conversation so far:
{history_text}

Follow-up question: {query}

Standalone question:"""

    for attempt in range(max_retries):
        try:
            result = model.invoke([HumanMessage(content=prompt)])
            rewritten = extract_text_from_gemini_response(result).strip()
            return rewritten if rewritten else query
        except Exception as e:
            error_text = str(e)
            if ("429" in error_text or "RESOURCE_EXHAUSTED" in error_text) and attempt < max_retries - 1:
                time.sleep(3 * (attempt + 1))
            else:
                return query  # Fall back to the original question
    return query


def retrieval_settings(total_chunks):
    """How many chunks each search checks (k_per_query) and how many go to the
    model (final_k), scaled to the size of THIS user's collection. Same
    thresholds answer_question always used; pulled out so the Settings panel
    can show the real numbers."""
    if total_chunks <= 100:
        k_per_query, final_k = 8, 10
    elif total_chunks <= 300:
        k_per_query, final_k = 10, 12
    elif total_chunks <= 600:
        k_per_query, final_k = 12, 15
    else:
        k_per_query, final_k = 15, 18
    return min(k_per_query, max(total_chunks, 1)), final_k


def answer_question(vectorstore, query, user_id, chat_history=None, k_per_query=6, final_k=8,
                    final_k_override=None, on_stage=None):
    """Retrieval + generation for one chat message.

    Pipeline: multi-query retrieval (original question + a few Gemini-generated
    rephrasings, each searched against ChromaDB — filtered to this user's own
    documents only — and deduplicated) -> BM25 reranking against the original
    question (narrows the combined candidate pool down to the most relevant
    chunks) -> those final chunks go to Gemini to generate the answer.

    This replaces the single-query MMR search used previously. Trade-off worth
    knowing: this does more work per question (one extra Gemini call to
    generate variations, plus up to ~5 similarity searches instead of 1), so
    each answer takes a bit longer and uses more of your API quota than
    before — in exchange for meaningfully better recall on questions where
    the exact wording doesn't closely match the source text.

    user_id (multi-user isolation): required. Every search below is scoped
    to only this user's own chunks, so this answer can never be grounded in
    another user's uploaded documents, however similar the topic.

    chat_history (history-aware generation): if the question is a follow-up
    that depends on earlier context ("what about the second one?"), it's
    rewritten into a standalone question first, using recent chat history.
    That standalone version is used for both retrieval and the final answer,
    so follow-ups actually find the right chunks instead of searching for
    the literal (context-free) words "the second one"."""
    if on_stage:
        on_stage("Understanding your question...")
    standalone_query = condense_question_with_history(query, chat_history)

    total_chunks = count_user_chunks(vectorstore, user_id)

    # Scale search breadth to the size of THIS USER's document collection
    # (not the shared total across all users). A fixed small k_per_query
    # works fine for one document, but becomes a real recall problem once
    # someone has uploaded many — each query variant was only checking 6
    # chunks out of a now much larger, more diverse corpus, so the right
    # chunk could easily not make the cut before reranking even saw it.
    k_per_query, final_k = retrieval_settings(total_chunks)
    if final_k_override:
        final_k = int(final_k_override)

    if on_stage:
        on_stage("Searching your documents...")
    candidates, queries_used = retrieve_with_multi_query(vectorstore, standalone_query, user_id, k_per_query=k_per_query)
    if on_stage:
        on_stage("Ranking the most relevant passages...")
    relevant_docs = rerank_with_bm25(standalone_query, candidates, top_n=min(final_k, len(candidates)))

    combined_input = f"""Based on the following documents, please answer this question: {standalone_query}

Documents:
{chr(10).join([f"- [{doc.metadata.get('source', 'unknown')}] {doc.page_content}" for doc in relevant_docs])}

Please provide a clear, helpful answer using only the information from these documents. If information from multiple documents is relevant, address both. If you can't find the answer in the documents, say "I don't have enough information to answer that question based on the provided documents."
"""

    if on_stage:
        on_stage("Writing the answer...")
    model = get_chat_model()
    messages = [
        SystemMessage(content="You are a helpful assistant."),
        HumanMessage(content=combined_input),
    ]
    result = model.invoke(messages)

    answer_text = extract_text_from_gemini_response(result)

    return answer_text, relevant_docs, standalone_query



# ---------- Password authentication ----------
# Accounts live in db/users.json as {user_id: {"salt": ..., "hash": ...}}.
# Passwords are never stored in plain text: each one is salted and hashed
# with PBKDF2-SHA256 (200k iterations), and checked with a constant-time
# comparison.
USERS_FILE = os.path.join("db", "users.json")
MIN_PASSWORD_LENGTH = 6


def load_users():
    if not os.path.exists(USERS_FILE):
        return {}
    try:
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_users(users):
    os.makedirs(os.path.dirname(USERS_FILE), exist_ok=True)
    tmp_path = USERS_FILE + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(users, f)
    os.replace(tmp_path, USERS_FILE)  # atomic write


def hash_password(password, salt_hex):
    return hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), 200_000
    ).hex()


def create_account(user_id, password):
    users = load_users()
    salt_hex = secrets.token_hex(16)
    users[user_id] = {"salt": salt_hex, "hash": hash_password(password, salt_hex)}
    save_users(users)


def verify_password(user_id, password):
    record = load_users().get(user_id)
    if not record:
        return False
    candidate = hash_password(password, record["salt"])
    return hmac.compare_digest(candidate, record["hash"])



def _add_markdown_text(paragraph, text):
    """Add text to a paragraph, turning **bold** and *italic* into real Word formatting."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)


def build_chat_docx(messages, user_id):
    """Build a Word (.docx) transcript of the chat and return it as bytes."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading("RAG Chat Transcript", level=1)
    info = doc.add_paragraph()
    info.add_run(f"User: {user_id}\n").italic = True
    info.add_run(f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True

    for m in messages:
        is_user = m["role"] == "user"
        label = doc.add_paragraph()
        label.paragraph_format.space_before = Pt(12)
        label.paragraph_format.space_after = Pt(2)
        run = label.add_run("You" if is_user else "Assistant")
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) if is_user else RGBColor(0x2E, 0x7D, 0x32)

        for line in m["content"].split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            heading = re.match(r"^#{1,6}\s+(.*)", stripped)
            bullet = re.match(r"^[-*\u2022]\s+(.*)", stripped)
            numbered = re.match(r"^\d+[.)]\s+(.*)", stripped)
            if heading:
                _add_markdown_text(doc.add_paragraph(), f"**{heading.group(1)}**")
            elif bullet:
                _add_markdown_text(doc.add_paragraph(style="List Bullet"), bullet.group(1))
            elif numbered:
                _add_markdown_text(doc.add_paragraph(style="List Number"), numbered.group(1))
            else:
                _add_markdown_text(doc.add_paragraph(), stripped)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------- Small talk (polite replies without searching documents) ----------
# Short social messages like "hi" or "thank you" get an instant, friendly
# reply instead of going through retrieval (which would search the documents
# for the word "hi" and answer "I don't have enough information"). This is
# rule-based, so it costs no API calls. It only triggers when the WHOLE
# message is small talk, so "hi, what is the refund policy?" still goes to
# the normal document Q&A.
SMALLTALK_RULES = [
    (
        r"(h+i+|h+e+y+|h+e+l+o+|hola|namaste|yo|sup|good (morning|afternoon|evening)|greetings)"
        r"( there| everyone| all| buddy| friend)?",
        ["Hi! 👋 How can I help you with your documents today?",
         "Hello! 😊 What would you like to know about your documents?",
         "Hey there! Ask me anything about your uploaded documents."],
    ),
    (
        r"((ok|okay|great|cool|nice) )?(thanks?|thank you|thank u|thx|thnx|ty|tysm|many thanks)"
        r"( (so|very) much| a lot| a ton| again| bro| buddy)*",
        ["You're welcome! 😊 Let me know if you need anything else.",
         "Welcome! Happy to help. Ask me anything else anytime.",
         "No problem at all! 👍"],
    ),
    (
        r"(bye+|goodbye|good night|see you|see ya|cya|take care|talk to you later|ttyl)( later| soon)?",
        ["Goodbye! 👋 Come back anytime.",
         "Take care! Your documents will be here when you return.",
         "Bye! Have a great day! 😊"],
    ),
    (
        r"(how are you|how are u|how r u|how do you do|how is it going|hows it going|"
        r"whats up|what is up|wassup)( today| doing)?",
        ["I'm doing great, thank you for asking! 😊 How can I help you today?",
         "All good here, thanks! What would you like to know about your documents?"],
    ),
    (
        r"(who are you|what are you|what can you do|what do you do|help|help me)",
        ["I'm a document assistant. Upload your files (PDF, Word, text or images) "
         "and ask me questions. I'll answer using only what's in them and show you the sources."],
    ),
    (
        r"(ok+|okay|k|cool|great|nice|awesome|got it|alright|fine|perfect|good|sure)",
        ["👍 Anything else you'd like to know?",
         "Great! Let me know if you have more questions."],
    ),
    (
        r"(sorry|my bad|oops|excuse me)",
        ["No problem at all! 😊 How can I help?"],
    ),
]


def get_smalltalk_reply(query):
    """Return a polite canned reply if the whole message is small talk, else None."""
    cleaned = re.sub(r"[^a-z0-9\s]", "", query.lower().replace("'", ""))
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if not cleaned or len(cleaned) > 40:
        return None
    for pattern, replies in SMALLTALK_RULES:
        if re.fullmatch(pattern, cleaned):
            return random.choice(replies)
    return None


# ---------- Streamlit UI ----------
# Everything above this line is the RAG engine (loading, chunking, embedding,
# retrieval, answering). Everything below is the interface. The UI only calls
# the engine's functions; it never changes how documents are processed.
#
# Design notes:
#  * Layout and controls use Streamlit's own widgets, so they keep working
#    when Streamlit changes its internals.
#  * Custom CSS is cosmetic polish only. It is scoped with `st.container(key=...)`
#    (a documented styling hook) plus a few stable data-testid attributes. If a
#    selector ever stops matching, the app loses some polish but keeps working.
#  * All colors adapt to light and dark themes (translucent surfaces, no
#    hard-coded backgrounds).

import streamlit.components.v1 as components

APP_NAME = "RAG Chat Agent"
ACCEPTED_TYPES = ["pdf", "docx", "txt", "jpg", "jpeg", "png", "webp", "bmp"]
IMAGE_TYPES = {"jpg", "jpeg", "png", "webp", "bmp"}
SUGGESTED_PROMPTS = [
    (":material/summarize:", "Summarize this document."),
    (":material/lightbulb:", "What are the key findings?"),
    (":material/school:", "Explain the main concepts."),
    (":material/manage_search:", "Find specific information in my documents."),
]
AVATARS = {"user": ":material/person:", "assistant": ":material/auto_awesome:"}
SOURCE_EXCERPT_CHARS = 700
CONTEXT_CHUNK_CHOICES = ["Auto", 6, 8, 10, 12, 15, 18]

st.set_page_config(
    page_title=APP_NAME,
    page_icon=":material/description:",
    layout="centered",
    initial_sidebar_state="auto",
)

logger = logging.getLogger("rag_chat")


# =====================================================================
# Design system: CSS
# =====================================================================
APP_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

:root { --rc-accent:#0F766E; --rc-accent-soft:rgba(15,118,110,.10);
        --rc-border:rgba(127,127,127,.26); --rc-surface:rgba(127,127,127,.07);
        --rc-radius:14px; --rc-radius-sm:10px; }

.stApp { font-family:'Inter',system-ui,-apple-system,'Segoe UI',Roboto,sans-serif;
         -webkit-font-smoothing:antialiased; }
[data-testid="stMainBlockContainer"], .block-container { max-width:880px; padding-top:2rem; padding-bottom:7rem; }
[data-testid="stSidebar"] { border-right:1px solid var(--rc-border); }

/* ---------- brand ---------- */
.rc-brand { display:flex; align-items:center; gap:.7rem; margin:.1rem 0 1rem; }
.rc-brand svg { width:34px; height:34px; flex:0 0 auto; }
.rc-brand-name { font-weight:650; font-size:1.02rem; letter-spacing:-.01em; line-height:1.2; }
.rc-brand-sub { font-size:.72rem; opacity:.6; }
.rc-section { font-size:.7rem; letter-spacing:.07em; text-transform:uppercase; font-weight:600;
              opacity:.55; margin:1.15rem 0 .4rem; }
.rc-summary { font-size:.8rem; opacity:.75; margin:0 0 .55rem; }
.rc-note { font-size:.74rem; opacity:.6; line-height:1.45; margin:.35rem 0; }

/* ---------- landing ---------- */
.rc-hero { text-align:center; padding:2.2rem 0 1.4rem; }
.rc-hero-logo svg { width:54px; height:54px; }
.rc-hero-title { font-size:2.15rem; font-weight:650; letter-spacing:-.03em; margin:.85rem 0 .45rem; line-height:1.15; }
.rc-hero-sub { max-width:34rem; margin:0 auto; opacity:.72; font-size:1.02rem; line-height:1.55; }
.rc-chips { display:flex; flex-wrap:wrap; gap:.4rem; justify-content:center; margin:1rem 0 .2rem; }
.rc-doc-chip { display:inline-flex; align-items:center; gap:.4rem; padding:.22rem .65rem; border-radius:999px;
               border:1px solid var(--rc-border); background:var(--rc-surface); font-size:.78rem; max-width:100%; }
.rc-doc-chip svg { width:13px; height:13px; flex:0 0 auto; opacity:.7; }
.rc-doc-chip span { overflow:hidden; text-overflow:ellipsis; white-space:nowrap; max-width:16rem; }
.rc-label { font-size:.78rem; font-weight:600; opacity:.65; margin:1.1rem 0 .5rem; text-align:center; }

/* suggestion cards */
.st-key-suggestions button { min-height:3.5rem; height:auto; text-align:left; justify-content:flex-start;
                             padding:.8rem 1rem; border-radius:12px; transition:border-color .15s, background .15s; }
.st-key-suggestions button:hover:not(:disabled) { border-color:var(--rc-accent); background:var(--rc-accent-soft); }

/* ---------- chat ---------- */
.rc-chat-title { font-weight:600; font-size:1.02rem; letter-spacing:-.01em; overflow:hidden;
                 text-overflow:ellipsis; white-space:nowrap; }
.rc-chat-sub { font-size:.74rem; opacity:.55; margin-top:.1rem; }
.rc-role { font-size:.76rem; font-weight:600; opacity:.65; margin-bottom:.2rem; }
.rc-meta { font-size:.72rem; opacity:.5; margin-top:.4rem; }

[data-testid="stChatMessage"] { border-radius:var(--rc-radius); padding:.9rem 1rem; margin-bottom:.4rem;
                                animation:rc-in .28s ease-out both; }
[data-testid="stChatMessage"]:has([data-testid="stChatMessageAvatarUser"]) { background:var(--rc-surface);
                                                                            border:1px solid var(--rc-border); }

/* typing + skeleton */
.rc-typing { display:flex; align-items:center; gap:.65rem; font-size:.9rem; opacity:.85; padding:.15rem 0 .5rem; }
.rc-dots span { display:inline-block; width:6px; height:6px; margin-right:3px; border-radius:50%;
                background:var(--rc-accent); animation:rc-bounce 1.2s infinite ease-in-out; }
.rc-dots span:nth-child(2) { animation-delay:.15s; } .rc-dots span:nth-child(3) { animation-delay:.3s; }
.rc-skel div { height:11px; border-radius:6px; margin:.55rem 0; background:linear-gradient(90deg,
               rgba(127,127,127,.14) 25%, rgba(127,127,127,.26) 37%, rgba(127,127,127,.14) 63%);
               background-size:400% 100%; animation:rc-shimmer 1.4s ease infinite; }
.rc-skel div:nth-child(1) { width:94%; } .rc-skel div:nth-child(2) { width:78%; } .rc-skel div:nth-child(3) { width:52%; }

/* error card */
.rc-error { display:flex; gap:.7rem; padding:.8rem .95rem; border-radius:var(--rc-radius-sm);
            border:1px solid rgba(180,35,24,.38); background:rgba(180,35,24,.07); }
.rc-error svg { width:18px; height:18px; flex:0 0 auto; margin-top:.1rem; color:#B42318;
                color:color-mix(in srgb,#B42318 75%,currentColor); }
.rc-error-title { font-weight:600; font-size:.92rem; }
.rc-error-hint { font-size:.84rem; opacity:.8; margin-top:.15rem; line-height:1.45; }

/* ---------- sources ---------- */
.rc-evidence-note { font-size:.78rem; opacity:.7; line-height:1.5; margin:.1rem 0 .6rem; }
.rc-src { border:1px solid var(--rc-border); border-radius:var(--rc-radius-sm); padding:.75rem .9rem;
          margin:.55rem 0; background:var(--rc-surface); transition:border-color .15s; }
.rc-src:hover { border-color:var(--rc-accent); }
.rc-src-head { display:flex; align-items:center; gap:.5rem; flex-wrap:wrap; }
.rc-src-num { width:21px; height:21px; border-radius:6px; display:grid; place-items:center; font-size:.7rem;
              font-weight:700; background:var(--rc-accent-soft); color:#0F766E;
              color:color-mix(in srgb,#0F766E 70%,currentColor); flex:0 0 auto; }
.rc-src-name { font-weight:600; font-size:.86rem; word-break:break-word; }
.rc-chip { font-size:.7rem; padding:.05rem .45rem; border:1px solid var(--rc-border); border-radius:6px; opacity:.85; }
.rc-quote { margin:.6rem 0 0; padding:.1rem 0 .1rem .8rem; border-left:3px solid var(--rc-accent); font-size:.84rem;
            line-height:1.6; opacity:.9; white-space:pre-wrap; word-break:break-word; }
.rc-src-foot { font-size:.7rem; opacity:.5; margin-top:.45rem; word-break:break-all; }

/* ---------- document cards ---------- */
.rc-file { display:flex; gap:.7rem; align-items:flex-start; padding:.65rem .75rem; border-radius:var(--rc-radius-sm);
           border:1px solid var(--rc-border); background:var(--rc-surface); transition:border-color .15s; }
.rc-file:hover { border-color:var(--rc-accent); }
.rc-file-icon { flex:0 0 34px; height:34px; border-radius:9px; display:grid; place-items:center;
                background:var(--rc-accent-soft); color:#0F766E; color:color-mix(in srgb,#0F766E 70%,currentColor); }
.rc-file-icon svg { width:18px; height:18px; }
.rc-file-body { min-width:0; flex:1; }
.rc-file-name { font-weight:560; font-size:.86rem; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.rc-file-meta { display:flex; flex-wrap:wrap; gap:.35rem; align-items:center; margin-top:.25rem; font-size:.72rem; opacity:.9; }
.rc-file-detail { font-size:.74rem; opacity:.75; margin-top:.35rem; line-height:1.4; }
.rc-pill { display:inline-flex; align-items:center; gap:.3rem; padding:.07rem .5rem; border-radius:999px;
           font-size:.7rem; font-weight:600; }
.rc-pill svg { width:12px; height:12px; }
.rc-pill-ready { background:rgba(15,118,110,.14); color:#0F766E; color:color-mix(in srgb,#0F766E 70%,currentColor); }
.rc-pill-busy  { background:rgba(180,83,9,.14);  color:#B45309; color:color-mix(in srgb,#B45309 70%,currentColor); }
.rc-pill-fail  { background:rgba(180,35,24,.13); color:#B42318; color:color-mix(in srgb,#B42318 70%,currentColor); }
.rc-pill-new   { background:rgba(127,127,127,.16); }
.rc-spin { animation:rc-spin 1s linear infinite; }

/* sidebar conversation list */
.st-key-chat_list button { justify-content:flex-start; text-align:left; }
.st-key-chat_list button p { overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
.st-key-conv_active button { background:var(--rc-accent-soft); }

/* general control polish */
.stButton > button, [data-testid="stPopover"] > button, [data-testid="stDownloadButton"] button
    { border-radius:10px; transition:border-color .15s, background .15s; }
.stButton > button:hover:not(:disabled) { border-color:var(--rc-accent); }

/* login */
.rc-login-title { text-align:center; font-size:1.55rem; font-weight:650; letter-spacing:-.025em; margin:.7rem 0 .25rem; }
.rc-login-sub { text-align:center; opacity:.7; font-size:.94rem; margin-bottom:1.1rem; line-height:1.5; }
.rc-login-logo { display:flex; justify-content:center; margin-top:2.2rem; }
.rc-login-logo svg { width:50px; height:50px; }

@keyframes rc-in { from { opacity:0; transform:translateY(6px); } to { opacity:1; transform:none; } }
@keyframes rc-bounce { 0%,80%,100% { transform:scale(.6); opacity:.5; } 40% { transform:scale(1); opacity:1; } }
@keyframes rc-shimmer { 0% { background-position:100% 50%; } 100% { background-position:0 50%; } }
@keyframes rc-spin { to { transform:rotate(360deg); } }

@media (max-width:640px) {
  [data-testid="stMainBlockContainer"], .block-container { padding-left:1rem; padding-right:1rem; padding-top:1.2rem; }
  .rc-hero { padding-top:1.2rem; } .rc-hero-title { font-size:1.65rem; } .rc-hero-sub { font-size:.95rem; }
  .rc-doc-chip span { max-width:11rem; }
}
@media (prefers-reduced-motion:reduce) {
  *, *::before, *::after { animation:none !important; transition:none !important; scroll-behavior:auto !important; }
}
</style>
"""


# =====================================================================
# Small HTML components (all user-provided text is escaped)
# =====================================================================
def esc(value):
    return html_lib.escape(str(value), quote=True)


ICON_PATHS = {
    "file": '<path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><polyline points="14 2 14 8 20 8"/>'
            '<line x1="16" y1="13" x2="8" y2="13"/><line x1="16" y1="17" x2="8" y2="17"/>',
    "image": '<rect x="3" y="3" width="18" height="18" rx="2" ry="2"/><circle cx="8.5" cy="8.5" r="1.5"/>'
             '<polyline points="21 15 16 10 5 21"/>',
    "check": '<polyline points="20 6 9 17 4 12"/>',
    "alert": '<circle cx="12" cy="12" r="10"/><line x1="12" y1="8" x2="12" y2="12"/><line x1="12" y1="16" x2="12.01" y2="16"/>',
    "loader": '<path d="M21 12a9 9 0 1 1-6.219-8.56"/>',
    "upload": '<polyline points="16 16 12 12 8 16"/><line x1="12" y1="12" x2="12" y2="21"/>'
              '<path d="M20.39 18.39A5 5 0 0 0 18 9h-1.26A8 8 0 1 0 3 16.3"/>',
}


def icon_svg(name, size=16, css_class=""):
    return (
        f'<svg class="{css_class}" width="{size}" height="{size}" viewBox="0 0 24 24" fill="none" '
        f'stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" '
        f'aria-hidden="true">{ICON_PATHS[name]}</svg>'
    )


def logo_svg(size=34):
    return (
        f'<svg width="{size}" height="{size}" viewBox="0 0 32 32" role="img" aria-label="{APP_NAME}">'
        '<rect width="32" height="32" rx="9" fill="#0F766E"/>'
        '<path d="M10.5 9h7l4 4v9.5a1.5 1.5 0 0 1-1.5 1.5h-9.5A1.5 1.5 0 0 1 9 22.5v-12A1.5 1.5 0 0 1 10.5 9z" '
        'fill="none" stroke="#fff" stroke-width="1.6" stroke-linejoin="round"/>'
        '<path d="M12.5 17.5h7M12.5 20.5h4.5" stroke="#fff" stroke-width="1.6" stroke-linecap="round"/>'
        '</svg>'
    )


def brand_html():
    return (f'<div class="rc-brand">{logo_svg(34)}<div><div class="rc-brand-name">{APP_NAME}</div>'
            '<div class="rc-brand-sub">Answers from your own documents</div></div></div>')


def section_html(title):
    return f'<div class="rc-section">{esc(title)}</div>'


def human_size(num_bytes):
    if not num_bytes:
        return ""
    size = float(num_bytes)
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
        size /= 1024


def file_ext(name):
    return os.path.splitext(name)[1].lower().lstrip(".")


def file_card_html(name, status, chunks=None, size=None, detail=None):
    """status: 'ready' (indexed, searchable), 'busy' (being indexed), 'fail'."""
    ext = file_ext(name)
    icon = icon_svg("image" if ext in IMAGE_TYPES else "file", 18)
    pill = {
        "ready": ("rc-pill-ready", icon_svg("check", 12), "Ready"),
        "busy": ("rc-pill-busy", icon_svg("loader", 12, "rc-spin"), "Indexing"),
        "fail": ("rc-pill-fail", icon_svg("alert", 12), "Failed"),
    }[status]
    meta = []
    if ext:
        meta.append(f'<span class="rc-chip">{esc(ext.upper())}</span>')
    if size:
        meta.append(f"<span>{esc(human_size(size))}</span>")
    if chunks:
        meta.append(f"<span>{chunks} chunk{'s' if chunks != 1 else ''}</span>")
    meta.append(f'<span class="rc-pill {pill[0]}">{pill[1]}{pill[2]}</span>')
    detail_html = f'<div class="rc-file-detail">{esc(detail)}</div>' if detail else ""
    return (
        f'<div class="rc-file"><div class="rc-file-icon">{icon}</div><div class="rc-file-body">'
        f'<div class="rc-file-name" title="{esc(name)}">{esc(name)}</div>'
        f'<div class="rc-file-meta">{"".join(meta)}</div>{detail_html}</div></div>'
    )


def typing_html(stage):
    return (
        f'<div class="rc-typing" role="status" aria-live="polite"><span class="rc-dots"><span></span><span></span>'
        f'<span></span></span><span>{esc(stage)}</span></div>'
        '<div class="rc-skel" aria-hidden="true"><div></div><div></div><div></div></div>'
    )


def error_card_html(title, hint):
    return (
        f'<div class="rc-error" role="alert">{icon_svg("alert", 18)}<div><div class="rc-error-title">{esc(title)}</div>'
        f'<div class="rc-error-hint">{esc(hint)}</div></div></div>'
    )


def source_card_html(index, src):
    chips = []
    page = src.get("page")
    if page is not None and str(page) != "":
        chips.append(f'<span class="rc-chip">Page {esc(page)}</span>')
    if src.get("section"):
        chips.append(f'<span class="rc-chip">Section: {esc(src["section"])}</span>')
    excerpt = src.get("excerpt", "")
    ellipsis = "…" if src.get("truncated") else ""
    foot = f'<div class="rc-src-foot">Chunk ID: {esc(src["chunk_id"])}</div>' if src.get("chunk_id") else ""
    return (
        f'<div class="rc-src"><div class="rc-src-head"><span class="rc-src-num">{index}</span>'
        f'<span class="rc-src-name">{esc(src.get("name", "unknown"))}</span>{"".join(chips)}</div>'
        f'<div class="rc-quote">{esc(excerpt)}{ellipsis}</div>{foot}</div>'
    )


# =====================================================================
# Friendly error messages (never expose raw exceptions or secrets)
# =====================================================================
def friendly_error(exc):
    """Map an exception to (title, hint) for display. The real exception is
    logged on the server only."""
    logger.exception("Answer generation failed")
    text = f"{type(exc).__name__} {exc}".lower()

    def has_code(*codes):
        return any(re.search(rf"\b{c}\b", text) for c in codes)

    if has_code("429") or "resource_exhausted" in text or "quota" in text or "rate limit" in text:
        return ("The AI service is busy or rate-limited.",
                "Wait a minute and try again. If this keeps happening, the API quota may be used up.")
    if "api key" in text or "api_key" in text or "permission_denied" in text or has_code("401", "403"):
        return ("The AI service rejected the request.",
                "The API key may be missing or invalid. The app owner should check the key configured for this app.")
    if has_code("404") or "not_found" in text or ("not found" in text and "model" in text):
        return ("The configured AI model isn't available.",
                "The app owner should check the model name in the app's settings.")
    if "timeout" in text or "timed out" in text or "deadline" in text or has_code("503", "504") \
            or "unavailable" in text or "overloaded" in text:
        return ("The AI service didn't respond in time.", "Try again in a moment.")
    if "connection" in text or "resolve" in text or "network" in text:
        return ("Couldn't reach the AI service.", "Check your connection and try again.")
    return ("Something went wrong while answering.",
            "Try again, or rephrase your question. If it keeps failing, remove and re-upload your document.")


def friendly_file_error(message):
    """Turn a raw ingestion error into a short, safe message for a file card."""
    text = str(message)
    low = text.lower()
    if low.startswith("unsupported file type") or low.startswith("no extractable text") \
            or low.startswith("could not read file with any known encoding"):
        return text.split(" Last error")[0][:200]
    if "429" in low or "resource_exhausted" in low or "quota" in low:
        return "The AI service was rate-limited while reading this file. Upload it again in a minute."
    if "timeout" in low or "deadline" in low or "unavailable" in low:
        return "The AI service timed out while reading this file. Upload it again."
    return "This file couldn't be processed. It may be corrupted, password-protected or empty."


# =====================================================================
# Conversations (kept in this browser session only)
# =====================================================================
def make_title(text):
    clean = " ".join(str(text).split())
    return clean if len(clean) <= 42 else clean[:40].rstrip() + "…"


def new_conversation():
    ss = st.session_state
    cid = uuid.uuid4().hex[:8]
    ss.conversations[cid] = {"id": cid, "title": "New chat", "messages": []}
    ss.conv_order.insert(0, cid)
    ss.active_conv = cid
    return cid


def active_conv():
    return st.session_state.conversations[st.session_state.active_conv]


def now_label():
    return datetime.now().strftime("%H:%M")


def doc_to_source(doc):
    """Only real metadata from the retrieval pipeline. Nothing is invented:
    there are no relevance scores because the pipeline doesn't return them."""
    meta = doc.metadata or {}
    text = doc.page_content or ""
    return {
        "name": os.path.basename(str(meta.get("source", "unknown")).replace("\\", "/")),
        "page": meta.get("page"),
        "section": meta.get("section"),
        "chunk_id": meta.get("chunk_id"),
        "excerpt": text[:SOURCE_EXCERPT_CHARS],
        "truncated": len(text) > SOURCE_EXCERPT_CHARS,
    }


def flash(kind, message):
    st.session_state.flash.append((kind, message))


# ---- the single entry point for "user asked something" ----
def dispatch_question(conv, text):
    """The user's message is already in conv. Reply to small talk instantly,
    explain if there are no documents yet, otherwise queue it for the RAG
    pipeline (which runs on the next rerun so the input can be disabled)."""
    ss = st.session_state
    reply = get_smalltalk_reply(text)
    if reply:
        conv["messages"].append({"role": "assistant", "content": reply, "time": now_label(), "kind": "smalltalk"})
        ss.scroll_pending = True
        return
    user_id = ss.authenticated_user
    if ss.vectorstore is None or count_user_chunks(ss.vectorstore, user_id) == 0:
        conv["messages"].append({
            "role": "assistant", "kind": "notice", "time": now_label(),
            "content": "You haven't added any documents yet. Upload a PDF, Word file, text file or image "
                       "and I can answer questions from it.",
        })
        ss.scroll_pending = True
        return
    ss.pending = {"conv": conv["id"], "query": text}
    ss.generating = True
    ss.scroll_pending = True


def queue_question(text):
    ss = st.session_state
    if ss.generating or not text.strip():
        return
    conv = active_conv()
    conv["messages"].append({"role": "user", "content": text.strip(), "time": now_label()})
    if len([m for m in conv["messages"] if m["role"] == "user"]) == 1:
        conv["title"] = make_title(text)
    dispatch_question(conv, text.strip())


# ---- callbacks (run before the script reruns) ----
def cb_new_chat():
    ss = st.session_state
    if ss.generating:
        return
    if not active_conv()["messages"]:
        return  # already on an empty chat
    new_conversation()


def cb_select_chat(cid):
    if not st.session_state.generating and cid in st.session_state.conversations:
        st.session_state.active_conv = cid


def cb_use_prompt(text):
    queue_question(text)


def cb_regenerate():
    ss = st.session_state
    conv = active_conv()
    if ss.generating:
        return
    msgs = conv["messages"]
    while msgs and msgs[-1]["role"] == "assistant":
        msgs.pop()
    if msgs and msgs[-1]["role"] == "user":
        dispatch_question(conv, msgs[-1]["content"])


def do_clear_chat(cid):
    conv = st.session_state.conversations.get(cid)
    if conv:
        conv["messages"] = []
        conv["title"] = "New chat"


def do_delete_chat(cid):
    ss = st.session_state
    ss.conversations.pop(cid, None)
    if cid in ss.conv_order:
        ss.conv_order.remove(cid)
    if not ss.conv_order:
        new_conversation()
    elif ss.active_conv == cid:
        ss.active_conv = ss.conv_order[0]


def do_delete_document(name):
    ss = st.session_state
    user_id = ss.authenticated_user
    try:
        delete_user_document(ss.vectorstore, user_id, name)
        ss.file_errors.pop(name, None)
        ss.file_sizes.pop(name, None)
        ss.ingested_filenames.discard(name)
        ss.uploader_version += 1
        flash("success", f"Removed {name}.")
    except Exception:
        logger.exception("Could not remove document")
        flash("error", f"Couldn't remove {name}. Please try again.")


def do_clear_documents():
    ss = st.session_state
    user_id = ss.authenticated_user
    try:
        # Deletes ONLY this user's chunks (filtered by user_id), never the
        # whole shared collection.
        if ss.vectorstore is not None:
            ss.vectorstore._collection.delete(where={"user_id": user_id})
        user_dir = os.path.join(DOCS_PATH, user_id)
        if os.path.isdir(user_dir):
            for leftover in os.listdir(user_dir):
                try:
                    os.remove(os.path.join(user_dir, leftover))
                except OSError:
                    pass
        ss.file_errors = {}
        ss.file_sizes = {}
        ss.ingested_filenames = set()
        ss.uploader_version += 1
        flash("success", "All your documents were removed.")
    except Exception:
        logger.exception("Could not clear documents")
        flash("error", "Couldn't remove your documents. Please try again.")


def cb_dismiss_error(name):
    ss = st.session_state
    ss.file_errors.pop(name, None)
    leftover = os.path.join(DOCS_PATH, ss.authenticated_user, name)
    if os.path.exists(leftover):
        try:
            os.remove(leftover)
        except OSError:
            pass


def cb_logout():
    ss = st.session_state
    for key in ("authenticated_user", "chat_owner", "conversations", "conv_order", "active_conv", "pending",
                "docx_cache"):
        ss.pop(key, None)
    ss.generating = False
    ss.file_errors = {}
    ss.file_sizes = {}
    ss.ingested_filenames = set()
    ss.uploader_version += 1


@st.dialog("Please confirm")
def confirm_dialog(message, confirm_label, action, args=()):
    st.write(message)
    left, right = st.columns(2)
    if left.button("Cancel", use_container_width=True):
        st.rerun()
    if right.button(confirm_label, type="primary", use_container_width=True):
        action(*args)
        st.rerun()


# =====================================================================
# Login
# =====================================================================
def normalize_user_id(raw):
    return raw.strip().lower().replace(" ", "_")


def render_login():
    try:
        url_user = st.query_params.get("user", "")
    except Exception:
        url_user = ""

    _, middle, _ = st.columns([1, 2.2, 1])
    with middle:
        st.markdown(
            f'<div class="rc-login-logo">{logo_svg(50)}</div>'
            f'<div class="rc-login-title">Welcome to {APP_NAME}</div>'
            '<div class="rc-login-sub">Upload your documents and ask questions about them. '
            'Sign in with a name and password to keep your documents private.</div>',
            unsafe_allow_html=True,
        )
        with st.container(border=True):
            raw = st.text_input(
                "Name or ID", value=url_user, key="login_name", placeholder="e.g. akhib, or your email",
                help="Use the same name each time to see your own documents again.",
            )
            uid = normalize_user_id(raw)
            if not uid:
                st.caption("Enter your name or ID, then press Enter.")
                return

            users = load_users()
            if uid not in users:
                st.caption(f"No account exists for **{uid}** yet. Choose a password to create one.")
                with st.form("signup_form"):
                    new_password = st.text_input("Choose a password", type="password")
                    confirm_password = st.text_input("Confirm password", type="password")
                    submitted = st.form_submit_button("Create account", type="primary", use_container_width=True)
                if submitted:
                    if len(new_password) < MIN_PASSWORD_LENGTH:
                        st.error(f"Password must be at least {MIN_PASSWORD_LENGTH} characters.")
                    elif new_password != confirm_password:
                        st.error("Passwords don't match.")
                    else:
                        create_account(uid, new_password)
                        st.session_state.authenticated_user = uid
                        st.rerun()
            else:
                with st.form("login_form"):
                    password = st.text_input("Password", type="password")
                    submitted = st.form_submit_button("Log in", type="primary", use_container_width=True)
                if submitted:
                    if verify_password(uid, password):
                        st.session_state.authenticated_user = uid
                        st.rerun()
                    else:
                        time.sleep(1)  # slows down password guessing
                        st.error("Incorrect password.")


# =====================================================================
# Document upload + indexing
# =====================================================================
def get_upload_limit_mb():
    try:
        return int(st.get_option("server.maxUploadSize"))
    except Exception:
        return 200


def render_uploader(label_visibility="collapsed"):
    return st.file_uploader(
        "Upload documents",
        type=ACCEPTED_TYPES,
        accept_multiple_files=True,
        key=f"uploader_{st.session_state.uploader_version}",
        label_visibility=label_visibility,
        help=f"PDF, Word (.docx), text, or images (JPG, PNG, WEBP, BMP). Up to {get_upload_limit_mb()} MB per file.",
    ) or []


def split_incoming(files, doc_files):
    """Separate files that need indexing from ones already indexed (so the same
    file is never embedded twice)."""
    to_process, already, seen = [], [], set()
    for f in files:
        if f.name in seen:
            continue
        seen.add(f.name)
        (already if f.name in doc_files else to_process).append(f)
    return to_process, already


def process_incoming(files, doc_files, user_id):
    """Index new files with the existing pipeline, showing real stages, then
    rerun. Distinguishes 'uploaded' from 'indexed and ready'."""
    ss = st.session_state
    to_process, already = split_incoming(files, doc_files)

    for f in already:
        flash("info", f"{f.name} is already indexed. Remove it first if you want to upload it again.")

    if not to_process:
        ss.uploader_version += 1
        st.rerun()

    for f in to_process:
        ss.file_errors.pop(f.name, None)
        ss.file_sizes[f.name] = getattr(f, "size", None)

    names = [f.name for f in to_process]
    status = st.status(f"Indexing {len(names)} document{'s' if len(names) != 1 else ''}", expanded=True)
    with status:
        st.markdown(
            "".join(
                f'<div class="rc-doc-chip" style="margin:.15rem .3rem .15rem 0">'
                f'<span class="rc-pill rc-pill-new">{icon_svg("upload", 12)}Uploaded</span>'
                f'<span>{esc(n)}</span></div>' for n in names
            ),
            unsafe_allow_html=True,
        )
        bar = st.progress(0.0, text="Preparing...")

    progress = {"frac": 0.0}

    def on_progress(current, total, status_text=None, **kwargs):
        text = status_text or kwargs.get("status")
        ratio = (current / total) if total else 0.0
        frac, label = progress["frac"], None
        if text:
            low = text.lower()
            label = text
            if low.startswith("saving"):
                frac, label = 0.02, "Saving uploaded files..."
            elif low.startswith("reading"):
                frac = 0.05 + 0.30 * ratio
            elif low.startswith("splitting"):
                frac, label = 0.38, "Splitting text into chunks..."
            elif low.startswith("embedding"):
                frac = 0.40 + 0.58 * ratio
        elif total == 0:
            label = "No content to embed."
        else:
            frac = 0.40 + 0.58 * ratio
            label = "Finishing up..." if current >= total else "Indexing chunks..."
        progress["frac"] = max(progress["frac"], min(frac, 1.0))
        bar.progress(progress["frac"], text=label or "Working...")

    try:
        ss.vectorstore = ingest_uploaded_files(to_process, user_id, progress_callback=on_progress)
        failed_hard = False
    except Exception as exc:
        failed_hard = True
        logger.exception("Document indexing failed")
        message = friendly_file_error(exc)
        for n in names:
            # remove any chunks stored before the failure so a half-indexed
            # file never shows up as "ready"
            try:
                delete_user_document(ss.vectorstore, user_id, n)
            except Exception:
                pass
            ss.file_errors[n] = message

    indexed_now = list_user_documents(ss.vectorstore, user_id)
    raw_errors = dict(ss.file_errors)  # raw messages recorded by the pipeline
    ready, failed = [], []
    for n in names:
        if not failed_hard and n in indexed_now:
            ready.append(n)
            ss.file_errors.pop(n, None)
        else:
            failed.append(n)
            if not failed_hard:
                ss.file_errors[n] = friendly_file_error(
                    raw_errors.get(n, "No extractable text found - it may be a scanned/image-only file or empty."))
    for n in names:
        ss.ingested_filenames.add(n)

    if ready:
        flash("success", f"{', '.join(ready)} {'is' if len(ready) == 1 else 'are'} indexed and ready for questions.")
    if failed:
        flash("error", f"Couldn't index: {', '.join(failed)}. See the document list for details.")

    status.update(label="Indexing finished" if ready else "Indexing failed",
                  state="complete" if ready else "error", expanded=False)
    ss.uploader_version += 1  # clear the dropzone so finished files aren't shown twice
    st.rerun()


def render_document_cards(doc_files, user_id, processing_names=()):
    ss = st.session_state
    for name in processing_names:
        st.markdown(file_card_html(name, "busy", size=ss.file_sizes.get(name)), unsafe_allow_html=True)

    for name, chunk_count in doc_files.items():
        size = ss.file_sizes.get(name)
        if not size:
            path = os.path.join(DOCS_PATH, user_id, name)
            size = os.path.getsize(path) if os.path.exists(path) else None
        card_col, action_col = st.columns([5, 1.1], vertical_alignment="center")
        card_col.markdown(file_card_html(name, "ready", chunks=chunk_count, size=size), unsafe_allow_html=True)
        if action_col.button(":material/delete:", key=f"rm_{name}", help=f"Remove {name}"):
            confirm_dialog(f"Remove **{name}** and its {chunk_count} indexed chunks? "
                           "Chats stay, but answers can no longer use this file.",
                           "Remove", do_delete_document, (name,))

    for name, message in list(ss.file_errors.items()):
        if name in doc_files or name in processing_names:
            continue
        card_col, action_col = st.columns([5, 1.1], vertical_alignment="center")
        card_col.markdown(file_card_html(name, "fail", detail=message, size=ss.file_sizes.get(name)),
                          unsafe_allow_html=True)
        action_col.button(":material/close:", key=f"dismiss_{name}", help="Dismiss. Upload the file again to retry.",
                          on_click=cb_dismiss_error, args=(name,))


# =====================================================================
# Sidebar
# =====================================================================
def settings_rows(total_chunks):
    rows = [("Answer model", CHAT_MODEL_NAME),
            ("Embedding model", EMBEDDING_MODEL_NAME.replace("models/", ""))]
    k_per_query, final_k = retrieval_settings(total_chunks)
    override = st.session_state.context_chunks
    rows.append(("Retrieval", "Multi-query search + keyword (BM25) re-rank"))
    rows.append(("Passages sent to model", f"{override if override != 'Auto' else final_k}"
                                           f"{' (auto)' if override == 'Auto' else ' (your setting)'}"))
    rows.append(("Chunks checked per search", str(k_per_query)))
    try:
        params = inspect.signature(ingestion.split_documents).parameters
        rows.append(("Chunking", f"Heading-aware, up to {params['chunk_size'].default:,} characters, "
                                 f"{params['chunk_overlap'].default} overlap"))
    except Exception:
        pass
    rows.append(("Scanned pages and images", "Read with Gemini vision"))
    return rows


def render_sidebar(user_id, doc_files, total_chunks, uploader_here):
    ss = st.session_state
    incoming = []
    with st.sidebar:
        st.markdown(brand_html(), unsafe_allow_html=True)
        st.button("New chat", icon=":material/add:", type="primary", use_container_width=True,
                  key="new_chat", on_click=cb_new_chat, disabled=ss.generating,
                  help="Start a fresh conversation")

        # ---- chats ----
        st.markdown(section_html("Chats"), unsafe_allow_html=True)
        order = ss.conv_order
        list_opts = {"key": "chat_list", "border": False}
        if len(order) > 6:
            list_opts["height"] = 260  # scroll box only when the list is long
        with st.container(**list_opts):
            for cid in order:
                conv = ss.conversations[cid]
                is_active = cid == ss.active_conv
                with st.container(key="conv_active" if is_active else f"conv_{cid}"):
                    st.button(conv["title"], key=f"sel_{cid}", use_container_width=True, type="tertiary",
                              icon=":material/chat_bubble:" if is_active else ":material/chat_bubble_outline:",
                              on_click=cb_select_chat, args=(cid,), disabled=ss.generating and not is_active)
        st.markdown('<div class="rc-note">Chats are kept for this browser session only. They are not saved '
                    'after you refresh or close the page.</div>', unsafe_allow_html=True)

        # ---- documents ----
        st.markdown(section_html("Documents"), unsafe_allow_html=True)
        if doc_files:
            st.markdown(f'<div class="rc-summary">{len(doc_files)} document{"s" if len(doc_files) != 1 else ""} '
                        f'indexed · {total_chunks} chunks searchable</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="rc-summary">No documents indexed yet.</div>', unsafe_allow_html=True)

        processing = []
        if uploader_here:
            incoming = render_uploader()
            st.markdown('<div class="rc-note">Scanned pages and images are read with Gemini vision, which takes '
                        'longer than plain text.</div>', unsafe_allow_html=True)
            to_process, _ = split_incoming(incoming, doc_files)
            processing = [f.name for f in to_process]
            for f in to_process:
                ss.file_sizes[f.name] = getattr(f, "size", None)

        render_document_cards(doc_files, user_id, processing_names=processing)

        if doc_files and st.button("Remove all documents", icon=":material/delete_forever:", key="rm_all",
                                   use_container_width=True):
            confirm_dialog("Remove **all** of your documents and their indexed chunks? This can't be undone.",
                           "Remove all", do_clear_documents)

        # ---- settings ----
        with st.expander("Settings & model info", icon=":material/tune:"):
            st.selectbox(
                "Passages sent to the model", CONTEXT_CHUNK_CHOICES, key="context_chunks",
                help="How many of the best-matching passages are given to the model when answering. "
                     "Auto scales with the size of your collection. More passages can help with broad "
                     "questions but make answers slower.",
            )
            rows = "".join(f'<tr><td style="opacity:.6;padding:.18rem .6rem .18rem 0;vertical-align:top;'
                           f'font-size:.76rem">{esc(k)}</td><td style="font-size:.76rem">{esc(v)}</td></tr>'
                           for k, v in settings_rows(total_chunks))
            st.markdown(f'<div class="rc-note"><table>{rows}</table></div>', unsafe_allow_html=True)

        # ---- account ----
        st.markdown(section_html("Account"), unsafe_allow_html=True)
        st.markdown(f'<div class="rc-summary">Signed in as <b>{esc(user_id)}</b></div>', unsafe_allow_html=True)
        st.button("Log out", icon=":material/logout:", key="logout", use_container_width=True, on_click=cb_logout,
                  disabled=ss.generating)
    return incoming


# =====================================================================
# Chat area
# =====================================================================
def docx_for_chat(conv, user_id):
    ss = st.session_state
    msgs = [{"role": m["role"], "content": m["content"]} for m in conv["messages"] if not m.get("error")]
    key = (conv["id"], len(msgs), hash(msgs[-1]["content"]) if msgs else 0)
    cached = ss.get("docx_cache")
    if cached and cached[0] == key:
        return cached[1]
    data = build_chat_docx(msgs, user_id)
    ss.docx_cache = (key, data)
    return data


def render_chat_header(conv, user_id, doc_count):
    left, right = st.columns([8, 1], vertical_alignment="center")
    left.markdown(
        f'<div class="rc-chat-title">{esc(conv["title"])}</div>'
        f'<div class="rc-chat-sub">{doc_count} document{"s" if doc_count != 1 else ""} available · '
        'session chat</div>', unsafe_allow_html=True)
    with right.popover(":material/more_horiz:", help="Chat options"):
        has_messages = bool(conv["messages"])
        st.download_button(
            "Download as Word", icon=":material/download:", use_container_width=True, disabled=not has_messages,
            data=docx_for_chat(conv, user_id) if has_messages else b"",
            file_name=f"chat_{user_id}_{datetime.now().strftime('%Y%m%d_%H%M')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key=f"dl_{conv['id']}",
        )
        if st.button("Clear messages", icon=":material/delete_sweep:", use_container_width=True,
                     key=f"clear_{conv['id']}", disabled=st.session_state.generating):
            confirm_dialog("Clear all messages in this chat?", "Clear", do_clear_chat, (conv["id"],))
        if st.button("Delete chat", icon=":material/delete:", use_container_width=True, key=f"del_{conv['id']}",
                     disabled=st.session_state.generating):
            confirm_dialog("Delete this chat?", "Delete", do_delete_chat, (conv["id"],))


def render_sources(msg):
    sources = msg.get("sources") or []
    if not sources:
        return
    files = {s["name"] for s in sources}
    label = f"Sources · {len(sources)} passage{'s' if len(sources) != 1 else ''} from " \
            f"{len(files)} document{'s' if len(files) != 1 else ''}"
    with st.expander(label, icon=":material/menu_book:"):
        st.markdown(
            '<div class="rc-evidence-note">These are the passages retrieved from your documents that the '
            'answer was based on, shown as stored. The answer above is written by the AI from them. '
            'Check the passages when accuracy matters.</div>', unsafe_allow_html=True)
        understood = msg.get("understood_as")
        if understood:
            st.markdown(f'<div class="rc-evidence-note">Your question was understood as: '
                        f'<i>{esc(understood)}</i></div>', unsafe_allow_html=True)
        box_opts = {"border": False}
        if len(sources) > 3:
            box_opts["height"] = 460  # scroll box only when there are many passages
        with st.container(**box_opts):
            for i, src in enumerate(sources, 1):
                st.markdown(source_card_html(i, src), unsafe_allow_html=True)


def render_message(conv, idx, msg, is_last):
    ss = st.session_state
    role = msg["role"]
    with st.chat_message(role, avatar=AVATARS[role]):
        st.markdown(f'<div class="rc-role">{"You" if role == "user" else APP_NAME}</div>', unsafe_allow_html=True)
        if msg.get("error"):
            st.markdown(error_card_html(msg["content"], msg.get("hint", "")), unsafe_allow_html=True)
        else:
            st.markdown(msg["content"])

        if role != "assistant":
            return
        if msg.get("time") and not msg.get("error"):
            st.markdown(f'<div class="rc-meta">{esc(msg["time"])}</div>', unsafe_allow_html=True)
        render_sources(msg)

        is_answer = msg.get("kind") not in ("smalltalk", "notice")
        if not is_answer:
            return
        copy_col, retry_col, _ = st.columns([1.2, 1.5, 4], gap="small", vertical_alignment="center")
        if not msg.get("error"):
            with copy_col.popover("Copy", icon=":material/content_copy:", help="Copy this response"):
                st.caption("Use the copy icon at the top-right of the box.")
                try:
                    st.code(msg["content"], language=None, wrap_lines=True)
                except TypeError:
                    st.code(msg["content"], language=None)
        if is_last:
            retry_col.button("Try again" if msg.get("error") else "Regenerate",
                             icon=":material/refresh:", key=f"regen_{conv['id']}_{idx}", on_click=cb_regenerate,
                             disabled=ss.generating, help="Ask the same question again")


def scroll_to_latest():
    components.html(
        "<script>try{const d=window.parent.document;const m=d.querySelectorAll('[data-testid=\"stChatMessage\"]');"
        "if(m.length){m[m.length-1].scrollIntoView({behavior:'smooth',block:'start'});}}catch(e){}</script>",
        height=0,
    )


def run_generation(user_id):
    """Answer the queued question using the existing RAG pipeline, with live
    stage updates, friendly errors, and the real sources saved on the message."""
    ss = st.session_state
    pending = ss.pending
    conv = ss.conversations.get(pending["conv"]) if pending else None
    if conv is None:
        ss.generating, ss.pending = False, None
        return

    query = pending["query"]
    history = [{"role": m["role"], "content": m["content"]} for m in conv["messages"][:-1]
               if not m.get("error") and m.get("kind") not in ("notice",)]
    override = ss.context_chunks if ss.context_chunks != "Auto" else None

    with st.chat_message("assistant", avatar=AVATARS["assistant"]):
        st.markdown(f'<div class="rc-role">{APP_NAME}</div>', unsafe_allow_html=True)
        slot = st.empty()
        slot.markdown(typing_html("Understanding your question..."), unsafe_allow_html=True)
        scroll_to_latest()

        def on_stage(text):
            slot.markdown(typing_html(text), unsafe_allow_html=True)

        try:
            answer, relevant_docs, standalone_query = answer_question(
                ss.vectorstore, query, user_id, chat_history=history,
                final_k_override=override, on_stage=on_stage,
            )
            understood = standalone_query if standalone_query.strip().lower() != query.strip().lower() else None
            message = {"role": "assistant", "content": answer, "time": now_label(), "understood_as": understood,
                       "sources": [doc_to_source(d) for d in relevant_docs]}
        except Exception as exc:
            title, hint = friendly_error(exc)
            message = {"role": "assistant", "content": title, "hint": hint, "error": True,
                       "time": now_label(), "kind": "error"}
        finally:
            ss.generating, ss.pending = False, None

    conv["messages"].append(message)
    ss.scroll_pending = True
    st.rerun()


def render_landing(doc_files, user_id, uploader_in_main):
    has_docs = bool(doc_files)
    st.markdown(
        f'<div class="rc-hero"><div class="rc-hero-logo">{logo_svg(54)}</div>'
        '<div class="rc-hero-title" role="heading" aria-level="1">Chat with your documents</div>'
        '<div class="rc-hero-sub">Upload PDFs, Word files, text files or images, then ask questions about '
        'their content. Answers come only from your documents, and the sources are shown.</div></div>',
        unsafe_allow_html=True)

    if uploader_in_main:
        with st.container(border=True):
            st.markdown('<div class="rc-label" style="margin-top:.2rem">Add your first document</div>',
                        unsafe_allow_html=True)
            files = render_uploader()
            st.markdown('<div class="rc-note" style="text-align:center">Scanned pages and images are read with '
                        'Gemini vision, which takes a little longer.</div>', unsafe_allow_html=True)
        if files:
            process_incoming(files, doc_files, user_id)
    elif has_docs:
        names = list(doc_files)[:6]
        chips = "".join(f'<span class="rc-doc-chip">{icon_svg("image" if file_ext(n) in IMAGE_TYPES else "file", 13)}'
                        f'<span>{esc(n)}</span></span>' for n in names)
        more = f'<span class="rc-doc-chip"><span>+{len(doc_files) - 6} more</span></span>' if len(doc_files) > 6 else ""
        st.markdown(f'<div class="rc-chips">{chips}{more}</div>', unsafe_allow_html=True)

    st.markdown('<div class="rc-label">Try asking</div>', unsafe_allow_html=True)
    with st.container(key="suggestions"):
        cols = st.columns(2)
        for i, (icon, text) in enumerate(SUGGESTED_PROMPTS):
            cols[i % 2].button(text, icon=icon, key=f"suggest_{i}", use_container_width=True,
                               on_click=cb_use_prompt, args=(text,), disabled=not has_docs,
                               help=None if has_docs else "Add a document first")


# =====================================================================
# App entry
# =====================================================================
st.markdown(APP_CSS, unsafe_allow_html=True)

ss = st.session_state
if "vectorstore" not in ss:
    ss.vectorstore = load_existing_vectorstore()
for _key, _default in (
    ("ingested_filenames", set()), ("uploader_version", 0), ("file_errors", {}), ("file_sizes", {}),
    ("flash", []), ("generating", False), ("pending", None), ("context_chunks", "Auto"),
    ("scroll_pending", False),
):
    if _key not in ss:
        ss[_key] = _default

user_id = ss.get("authenticated_user")
if not user_id:
    render_login()
    st.stop()

# chats belong to the signed-in user for this browser session
if ss.get("chat_owner") != user_id or "conversations" not in ss:
    ss.conversations, ss.conv_order, ss.chat_owner = {}, [], user_id
    ss.generating, ss.pending = False, None
    new_conversation()
if ss.active_conv not in ss.conversations:
    ss.active_conv = ss.conv_order[0]

conv = active_conv()
doc_files = list_user_documents(ss.vectorstore, user_id)
total_chunks = sum(doc_files.values())
landing_mode = not conv["messages"] and not ss.generating
uploader_in_main = landing_mode and not doc_files and not ss.file_errors

sidebar_files = render_sidebar(user_id, doc_files, total_chunks, uploader_here=not uploader_in_main)

# one-time messages from the previous run (toasts for success, banners for problems)
notices, ss.flash = ss.flash, []
for kind, text in notices:
    if kind == "success":
        st.toast(text, icon=":material/check_circle:")
    elif kind == "info":
        st.toast(text, icon=":material/info:")
    elif kind == "warning":
        st.warning(text, icon=":material/warning:")
    else:
        st.error(text, icon=":material/error:")

if sidebar_files:
    process_incoming(sidebar_files, doc_files, user_id)

if landing_mode:
    render_landing(doc_files, user_id, uploader_in_main)
else:
    render_chat_header(conv, user_id, len(doc_files))
    for i, message in enumerate(conv["messages"]):
        render_message(conv, i, message, is_last=(i == len(conv["messages"]) - 1))

# ---- composer (docked to the bottom by Streamlit) ----
composer = {
    "placeholder": "Ask anything about your documents...",
    "key": "composer",
    "disabled": ss.generating,
}
if "accept_file" in inspect.signature(st.chat_input).parameters:
    composer.update(accept_file="multiple", file_type=ACCEPTED_TYPES)  # attachment button
submitted = st.chat_input(**composer)

if submitted:
    if isinstance(submitted, str):
        typed_text, attached = submitted, []
    else:
        typed_text = getattr(submitted, "text", None) or (submitted.get("text") if hasattr(submitted, "get") else "") or ""
        attached = list(getattr(submitted, "files", None) or (submitted.get("files") if hasattr(submitted, "get") else []) or [])
    if attached:
        if typed_text.strip():
            ss.deferred_question = typed_text.strip()
        process_incoming(attached, doc_files, user_id)  # ends with st.rerun()
    elif typed_text.strip():
        queue_question(typed_text)
        st.rerun()

# a question typed together with an attachment is asked once the files are indexed
if ss.get("deferred_question") and not ss.generating:
    deferred = ss.pop("deferred_question")
    queue_question(deferred)
    st.rerun()

if ss.generating and ss.pending:
    run_generation(user_id)

if ss.scroll_pending:
    ss.scroll_pending = False
    scroll_to_latest()
